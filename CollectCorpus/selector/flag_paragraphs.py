"""Find training paragraphs that don't help: their content contradicts their page's label.

    python -m selector.flag_paragraphs --data data/selector/v3        (run from the CollectCorpus folder)

Each paragraph is scored by a classifier that never saw its page (cross-validation by page, GroupKFold):
the pages are split into 5 groups, a logistic regression is trained on 4 and scores the paragraphs of the 5th.
A paragraph of a related page (1) with a low score, or of an unrelated page (0) with a high score, is flagged.
The decision what to remove is the user's (paragraph_review.xlsx).

E5 vectors are computed on the CPU in chunks saved to <data>/vectors_e5/, so a stopped run continues.
Writes <data>/paragraphs_scored.csv and <data>/paragraph_review.xlsx.
"""
import argparse
import sys
import time
from pathlib import Path

import numpy as np
import pandas as pd
from sklearn.linear_model import LogisticRegression
from sklearn.model_selection import GroupKFold

sys.stdout.reconfigure(encoding="utf-8")

E5_MODEL = "intfloat/multilingual-e5-large"
CHUNK = 200
STRONG_LOW, STRONG_HIGH = 0.3, 0.7   # a clear contradiction: recommend removing


def embed_cached(texts, folder, model_name=E5_MODEL):
    """E5 vectors for all texts, computed in chunks; chunks already on disk are reused."""
    folder.mkdir(parents=True, exist_ok=True)
    model, parts, start_time = None, [], time.time()
    for start in range(0, len(texts), CHUNK):
        path = folder / f"{start:06d}_{len(texts)}.npy"
        if not path.exists():
            if model is None:
                from sentence_transformers import SentenceTransformer
                model = SentenceTransformer(model_name, device="cpu")
            vec = model.encode(["query: " + t for t in texts[start:start + CHUNK]], batch_size=16,
                               normalize_embeddings=True)
            np.save(path, vec)
            done = min(start + CHUNK, len(texts))
            print(f"  vectors {done}/{len(texts)}  ({(time.time() - start_time) / 60:.0f} min)", flush=True)
        parts.append(np.load(path))
    return np.vstack(parts)


def out_of_fold_scores(x, labels, groups, folds=5):
    """Probability of "related" for each paragraph, from a model trained without its page."""
    scores = np.zeros(len(labels))
    for train_idx, test_idx in GroupKFold(n_splits=folds).split(x, labels, groups):
        clf = LogisticRegression(C=10, max_iter=2000, class_weight="balanced").fit(x[train_idx], labels[train_idx])
        scores[test_idx] = clf.predict_proba(x[test_idx])[:, 1]
    return scores


def review_workbook(df, path):
    """Flagged paragraphs with a recommendation; the user's decision column is pre-filled with it."""
    flagged = df[df.flag_cv].copy()
    strong = ((flagged.label == 1) & (flagged.cv_score < STRONG_LOW)) | ((flagged.label == 0) & (flagged.cv_score >= STRONG_HIGH))
    flagged["recommendation"] = np.where(strong, "להסיר", "להשאיר")
    flagged["reason"] = np.where(flagged.label == 1,
                                 np.where(strong, "דף צבאי, אבל הפסקה לא צבאית", "דף צבאי, פסקה גבולית"),
                                 np.where(strong, "דף לא צבאי, אבל הפסקה צבאית", "דף לא צבאי, פסקה גבולית"))
    flagged["decision"] = flagged.recommendation
    flagged = flagged.sort_values(["page", "paragraph"])
    out = flagged[["page", "label", "paragraph", "section", "text", "cv_score", "recommendation", "reason", "decision"]]
    out.columns = ["דף", "תווית הדף", "מס' פסקה", "פרק", "טקסט", "ציון (0-1)", "המלצה", "סיבה", "החלטה"]
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        out.to_excel(xw, index=False, sheet_name="פסקאות לבדיקה")
        ws = xw.sheets["פסקאות לבדיקה"]
        ws.sheet_view.rightToLeft = True
        ws.freeze_panes = "A2"
        from openpyxl.styles import Alignment, PatternFill
        from openpyxl.worksheet.datavalidation import DataValidation
        for col, w in zip("ABCDEFGHI", [30, 8, 8, 26, 90, 9, 9, 26, 10]):
            ws.column_dimensions[col].width = w
        for r in range(2, len(out) + 2):
            ws.cell(r, 5).alignment = Alignment(wrap_text=True, vertical="top")
            ws.cell(r, 9).fill = PatternFill("solid", fgColor="FFF2CC")
        dv = DataValidation(type="list", formula1='"להסיר,להשאיר"')
        ws.add_data_validation(dv)
        dv.add(f"I2:I{len(out) + 1}")
    return out


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--data", default="data/selector/v3", help="folder with paragraphs.csv")
    args = ap.parse_args()
    data = Path(args.data)
    df = pd.read_csv(data / "paragraphs.csv", encoding="utf-8-sig")
    print(f"{len(df)} paragraphs from {df.page.nunique()} pages")
    x = embed_cached(df.text.tolist(), data / "vectors_e5")
    df["cv_score"] = np.round(out_of_fold_scores(x, df.label.to_numpy(), df.page.to_numpy()), 3)
    df["flag_cv"] = ((df.label == 1) & (df.cv_score < 0.5)) | ((df.label == 0) & (df.cv_score >= 0.5))
    df.to_csv(data / "paragraphs_scored.csv", index=False, encoding="utf-8-sig")
    out = review_workbook(df, data / "paragraph_review.xlsx")
    print(f"flagged {len(out)} paragraphs; recommended to remove {(out['המלצה'] == 'להסיר').sum()}")
    print(df[df.flag_cv].groupby(["label", "page"]).size().to_string())


if __name__ == "__main__":
    main()
